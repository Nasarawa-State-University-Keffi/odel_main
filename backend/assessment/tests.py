"""
Comprehensive tests for Moodle-style Quiz System

Tests cover:
- Question type plugins (multichoice, truefalse, shortanswer, essay)
- Question and Quiz services
- REST API endpoints
- Grading logic and partial credit
- Multiple attempts
- Manual grading

CORRECTED VERSION - All tests should pass
"""
import uuid
from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from openpyxl import load_workbook
from django.test import TestCase
from django.utils import timezone
from django.core.exceptions import ValidationError
from rest_framework.test import APITestCase, APIClient
from rest_framework import status
from portal_auth.models import PortalUser
from courses.models import (
    AcademicSession,
    CourseCache,
    CourseOffering,
    Semester,
    StudentRegisteredCourse,
    StaffAssignedCourse,
)
from .models import (
    QuestionCategory, Question, QuestionAnswer,
    Quiz, QuizQuestion, QuizAttempt, QuestionAttempt,
    Assignment, AssignmentSubmission,
    Assessment, AssessmentQuestion, AssessmentAttempt, AssessmentQuestionAttempt, Grade,
)
from .services import AssessmentService, QuizService, QuestionService, calculate_student_total, get_gradebook_summary
from .question_types import get_question_type_handler


class QuestionTypePluginTests(TestCase):
    """Test question type plugin system"""
    
    def setUp(self):
        """Create test course and category"""
        self.course = CourseCache.objects.create(
            course_external_id=101,
            course_title='Test Course',
            course_code='TEST101'
        )
        self.category = QuestionCategory.objects.create(
            course=self.course,
            name='Test Category'
        )
    
    def test_multiple_choice_single_correct(self):
        """Test multiple choice with single correct answer"""
        question = Question.objects.create(
            category=self.category,
            qtype='multichoice',
            name='Python Variables',
            question_text='Which is a valid Python variable?',
            default_mark=Decimal('10.00')
        )
        
        # Create answers
        correct = QuestionAnswer.objects.create(
            question=question,
            answer_text='my_var',
            fraction=Decimal('1.0'),
            order=1
        )
        QuestionAnswer.objects.create(
            question=question,
            answer_text='2nd_var',
            fraction=Decimal('0.0'),
            order=2
        )
        
        handler = get_question_type_handler('multichoice')
        
        # Test correct answer
        response = {'selected': [str(correct.id)]}
        is_valid, _ = handler.validate_response(question, response)
        self.assertTrue(is_valid)
        
        fraction = handler.grade(question, response)
        self.assertEqual(fraction, Decimal('1.0'))
    
    def test_multiple_choice_partial_credit(self):
        """Test multiple choice with partial credit"""
        question = Question.objects.create(
            category=self.category,
            qtype='multichoice',
            name='Python Data Types',
            question_text='Which are mutable in Python?',
            default_mark=Decimal('10.00')
        )
        
        # Create answers with partial credit
        ans1 = QuestionAnswer.objects.create(
            question=question,
            answer_text='list',
            fraction=Decimal('0.5'),
            order=1
        )
        ans2 = QuestionAnswer.objects.create(
            question=question,
            answer_text='dict',
            fraction=Decimal('0.5'),
            order=2
        )
        QuestionAnswer.objects.create(
            question=question,
            answer_text='tuple',
            fraction=Decimal('0.0'),
            order=3
        )
        
        handler = get_question_type_handler('multichoice')
        
        # Test selecting one correct answer (50%)
        response = {'selected': [str(ans1.id)]}
        fraction = handler.grade(question, response)
        self.assertEqual(fraction, Decimal('0.5'))
        
        # Test selecting both correct answers (100%)
        response = {'selected': [str(ans1.id), str(ans2.id)]}
        fraction = handler.grade(question, response)
        self.assertEqual(fraction, Decimal('1.0'))

    def test_single_choice_accepts_one_option_and_rejects_multiple_answers(self):
        question = Question.objects.create(
            category=self.category,
            qtype='singlechoice',
            name='Single answer question',
            question_text='Which planet is known as the Red Planet?',
            default_mark=Decimal('5.00'),
        )
        correct = QuestionAnswer.objects.create(
            question=question, answer_text='Mars', fraction=Decimal('1.00'), order=1,
        )
        incorrect = QuestionAnswer.objects.create(
            question=question, answer_text='Venus', fraction=Decimal('0.00'), order=2,
        )
        handler = get_question_type_handler('singlechoice')

        valid, _ = handler.validate_response(question, {'selected': str(correct.id)})
        self.assertTrue(valid)
        self.assertEqual(handler.grade(question, {'selected': str(correct.id)}), Decimal('1.0'))
        self.assertEqual(handler.grade(question, {'selected': str(incorrect.id)}), Decimal('0.0'))

        valid, message = handler.validate_response(
            question, {'selected': [str(correct.id), str(incorrect.id)]},
        )
        self.assertFalse(valid)
        self.assertIn('one answer ID', message)
    
    def test_true_false_question(self):
        """Test true/false question type"""
        question = Question.objects.create(
            category=self.category,
            qtype='truefalse',
            name='Python Is Compiled',
            question_text='Python is a compiled language',
            default_mark=Decimal('5.00')
        )
        
        # Create True/False answers
        false_ans = QuestionAnswer.objects.create(
            question=question,
            answer_text='False',
            fraction=Decimal('1.0'),
            order=1
        )
        true_ans = QuestionAnswer.objects.create(
            question=question,
            answer_text='True',
            fraction=Decimal('0.0'),
            order=2
        )
        
        handler = get_question_type_handler('truefalse')
        
        # Test correct answer (False)
        response = {'selected': str(false_ans.id)}
        is_valid, _ = handler.validate_response(question, response)
        self.assertTrue(is_valid)
        fraction = handler.grade(question, response)
        self.assertEqual(fraction, Decimal('1.0'))
        
        # Test incorrect answer (True)
        response = {'selected': str(true_ans.id)}
        fraction = handler.grade(question, response)
        self.assertEqual(fraction, Decimal('0.0'))

    def test_true_false_rejects_invalid_scoring_configuration(self):
        question = Question.objects.create(
            category=self.category,
            qtype='truefalse',
            name='Invalid Boolean configuration',
            question_text='This must never create a negative score.',
            default_mark=Decimal('5.00'),
        )
        selected = QuestionAnswer.objects.create(
            question=question, answer_text='True', fraction=Decimal('-1.00'), order=1,
        )
        QuestionAnswer.objects.create(
            question=question, answer_text='False', fraction=Decimal('0.50'), order=2,
        )
        handler = get_question_type_handler('truefalse')
        valid, message = handler.validate_response(question, {'selected': str(selected.id)})
        self.assertFalse(valid)
        self.assertIn('one correct and one incorrect', message)
        self.assertEqual(handler.grade(question, {'selected': str(selected.id)}), Decimal('0.0'))
    
    def test_short_answer_case_insensitive(self):
        """Test short answer with case-insensitive matching"""
        question = Question.objects.create(
            category=self.category,
            qtype='shortanswer',
            name='Capital of France',
            question_text='What is the capital of France?',
            default_mark=Decimal('5.00')
        )
        
        # Create acceptable answers
        QuestionAnswer.objects.create(
            question=question,
            answer_text='Paris',
            fraction=Decimal('1.0'),
            order=1
        )
        QuestionAnswer.objects.create(
            question=question,
            answer_text='paris',
            fraction=Decimal('1.0'),
            order=2
        )
        
        handler = get_question_type_handler('shortanswer')
        
        # Test various case variations
        for answer in ['Paris', 'PARIS', 'paris', 'PaRiS']:
            response = {'text': answer}
            is_valid, _ = handler.validate_response(question, response)
            self.assertTrue(is_valid)
            fraction = handler.grade(question, response)
            self.assertEqual(fraction, Decimal('1.0'), f'Failed for: {answer}')
        
        # Test incorrect answer
        response = {'text': 'London'}
        fraction = handler.grade(question, response)
        self.assertEqual(fraction, Decimal('0.0'))
    
    def test_essay_requires_manual_grading(self):
        """Test essay question returns None for auto-grading"""
        question = Question.objects.create(
            category=self.category,
            qtype='essay',
            name='Python Benefits',
            question_text='Explain the benefits of Python',
            default_mark=Decimal('20.00')
        )
        
        handler = get_question_type_handler('essay')
        
        response = {'text': 'Python is great because...'}
        is_valid, _ = handler.validate_response(question, response)
        self.assertTrue(is_valid)
        
        # Essay should return None (requires manual grading)
        fraction = handler.grade(question, response)
        self.assertIsNone(fraction)


class QuestionServiceTests(TestCase):
    """Test QuestionService business logic"""
    
    def setUp(self):
        self.course = CourseCache.objects.create(
            course_external_id=101,
            course_title='Test Course',
            course_code='TEST101'
        )
        self.category = QuestionCategory.objects.create(
            course=self.course,
            name='Test Category'
        )
    
    def test_validate_question_type(self):
        """Test question type validation"""
        self.assertTrue(QuestionService.validate_question_type('multichoice'))
        self.assertTrue(QuestionService.validate_question_type('singlechoice'))
        self.assertTrue(QuestionService.validate_question_type('truefalse'))
        self.assertTrue(QuestionService.validate_question_type('shortanswer'))
        self.assertTrue(QuestionService.validate_question_type('essay'))
        self.assertFalse(QuestionService.validate_question_type('invalid'))
    
    def test_grade_question(self):
        """Test question grading through service"""
        question = Question.objects.create(
            category=self.category,
            qtype='shortanswer',
            name='Test Question',
            question_text='What is 2+2?',
            default_mark=Decimal('10.00')
        )
        
        QuestionAnswer.objects.create(
            question=question,
            answer_text='4',
            fraction=Decimal('1.0'),
            order=1
        )
        
        response = {'text': '4'}
        fraction, score = QuestionService.grade_question(question, response)
        
        self.assertEqual(fraction, Decimal('1.0'))
        self.assertEqual(score, Decimal('10.00'))
    
    def test_get_correct_answer(self):
        """Test retrieving correct answer"""
        question = Question.objects.create(
            category=self.category,
            qtype='multichoice',
            name='Test Question',
            question_text='Select correct',
            default_mark=Decimal('10.00')
        )
        
        correct = QuestionAnswer.objects.create(
            question=question,
            answer_text='Correct Answer',
            fraction=Decimal('1.0'),
            order=1
        )
        
        result = QuestionService.get_correct_answer(question)
        # The service returns 'answers' not 'correct_answers'
        self.assertIn('answers', result)
        self.assertEqual(len(result['correct_ids']), 1)
        # Find the correct answer in the answers list
        correct_answer = [a for a in result['answers'] if a['fraction'] == 1.0][0]
        self.assertEqual(correct_answer['text'], 'Correct Answer')


class QuizServiceTests(TestCase):
    """Test QuizService business logic"""
    
    def setUp(self):
        self.course = CourseCache.objects.create(
            course_external_id=101,
            course_title='Test Course',
            course_code='TEST101'
        )
        self.category = QuestionCategory.objects.create(
            course=self.course,
            name='Test Category'
        )
        
        # Create quiz
        self.quiz = Quiz.objects.create(
            course=self.course,
            name='Test Quiz',
            is_published=True,
            time_limit=3600,
            max_grade=Decimal('100.00'),
            max_attempts=3
        )
        
        # Create question
        self.question = Question.objects.create(
            category=self.category,
            qtype='shortanswer',
            name='Test Question',
            question_text='What is 2+2?',
            default_mark=Decimal('10.00')
        )
        
        QuestionAnswer.objects.create(
            question=self.question,
            answer_text='4',
            fraction=Decimal('1.0'),
            order=1
        )
        
        # Link question to quiz
        QuizQuestion.objects.create(
            quiz=self.quiz,
            question=self.question,
            order=1,
            max_mark=Decimal('10.00')
        )
    
    def test_start_attempt(self):
        """Test starting a quiz attempt"""
        attempt = QuizService.start_attempt(
            quiz=self.quiz,
            user_external_id='student_1'
        )
        
        self.assertIsNotNone(attempt)
        self.assertEqual(attempt.state, 'in_progress')
        self.assertEqual(attempt.attempt_number, 1)
        
        # Verify question attempts were created
        question_attempts = attempt.question_attempts.all()
        self.assertEqual(question_attempts.count(), 1)
    
    def test_max_attempts_enforcement(self):
        """Test that max attempts is enforced"""
        user_id = 'student@test.com'
        
        # Create and finish 3 attempts (must finish to allow new attempts)
        for i in range(3):
            attempt = QuizService.start_attempt(quiz=self.quiz, user_external_id=user_id)
            QuizService.finish_attempt(attempt)
        
        # 4th attempt should fail
        with self.assertRaises(ValidationError) as context:
            QuizService.start_attempt(quiz=self.quiz, user_external_id=user_id)
        
        self.assertIn('Maximum attempts', str(context.exception))
    
    def test_submit_response(self):
        """Test submitting a response"""
        attempt = QuizService.start_attempt(
            quiz=self.quiz,
            user_external_id='student_1'
        )
        
        response = {'text': '4'}
        question_attempt = QuizService.submit_response(
            attempt=attempt,
            question=self.question,
            response=response
        )
        
        self.assertIsNotNone(question_attempt)
        self.assertEqual(question_attempt.fraction, Decimal('1.0'))
        self.assertEqual(question_attempt.score, Decimal('10.00'))
        self.assertIsNotNone(question_attempt.graded_at)
    
    def test_finish_attempt(self):
        """Test finishing a quiz attempt"""
        attempt = QuizService.start_attempt(
            quiz=self.quiz,
            user_external_id='student_1'
        )
        
        # Submit response
        QuizService.submit_response(
            attempt=attempt,
            question=self.question,
            response={'text': '4'}
        )
        
        # Finish attempt
        QuizService.finish_attempt(attempt)
        
        attempt.refresh_from_db()
        self.assertEqual(attempt.state, 'finished')
        self.assertIsNotNone(attempt.finished_at)
        # Score: 10 out of 10 possible = 100%, scaled to max_grade of 100
        self.assertEqual(attempt.total_score, Decimal('100.00'))
    
    def test_calculate_final_grade(self):
        """Test final grade calculation"""
        attempt = QuizService.start_attempt(
            quiz=self.quiz,
            user_external_id='student_1'
        )
        
        # Submit correct response
        QuizService.submit_response(
            attempt=attempt,
            question=self.question,
            response={'text': '4'}
        )
        
        # calculate_final_grade returns the score but doesn't save it
        final_score = QuizService.calculate_final_grade(attempt)
        
        # Score: 10 out of 10 possible = 100%, scaled to max_grade of 100
        self.assertEqual(final_score, Decimal('100.00'))
    
    def test_get_attempt_summary(self):
        """Test attempt summary statistics"""
        attempt = QuizService.start_attempt(
            quiz=self.quiz,
            user_external_id='student_1'
        )
        
        QuizService.submit_response(
            attempt=attempt,
            question=self.question,
            response={'text': '4'}
        )
        
        summary = QuizService.get_attempt_summary(attempt)
        
        self.assertIn('total_questions', summary)
        self.assertIn('answered_questions', summary)
        self.assertIn('graded_questions', summary)
        self.assertIn('total_score', summary)
        self.assertEqual(summary['total_questions'], 1)
        self.assertEqual(summary['answered_questions'], 1)
        self.assertEqual(summary['graded_questions'], 1)


class StaffAssignmentContextAPITests(APITestCase):
    """Assignment lists must stay within the teaching period selected by staff."""

    def setUp(self):
        self.instructor = PortalUser.objects.create(
            external_id='assignment_instructor',
            full_name='Assignment Instructor',
            is_staff=True,
        )
        self.session = AcademicSession.objects.create(name='2025/2026')
        self.first_semester = Semester.objects.create(name='First Semester')
        self.second_semester = Semester.objects.create(name='Second Semester')
        self.first_course = CourseCache.objects.create(
            course_external_id=401,
            course_title='First Semester Course',
            course_code='FIRST401',
        )
        self.second_course = CourseCache.objects.create(
            course_external_id=402,
            course_title='Second Semester Course',
            course_code='SECOND402',
        )
        first_offering = CourseOffering.objects.create(
            course=self.first_course,
            session=self.session,
            semester=self.first_semester,
            programme_type_code='ODEL',
        )
        second_offering = CourseOffering.objects.create(
            course=self.second_course,
            session=self.session,
            semester=self.second_semester,
            programme_type_code='ODEL',
        )
        StaffAssignedCourse.objects.create(
            staff_external_id=self.instructor.external_id,
            course=self.first_course,
            course_offering=first_offering,
            programme_type_code='ODEL',
            role='INSTRUCTOR',
        )
        StaffAssignedCourse.objects.create(
            staff_external_id=self.instructor.external_id,
            course=self.second_course,
            course_offering=second_offering,
            programme_type_code='ODEL',
            role='INSTRUCTOR',
        )
        Assignment.objects.create(
            course=self.first_course,
            title='First semester assignment',
            open_at=timezone.now(),
            due_at=timezone.now() + timedelta(days=7),
            created_by=self.instructor,
        )
        Assignment.objects.create(
            course=self.second_course,
            title='Second semester assignment',
            open_at=timezone.now(),
            due_at=timezone.now() + timedelta(days=7),
            created_by=self.instructor,
        )
        self.client.force_authenticate(user=self.instructor)

    def test_assignment_list_requires_a_complete_teaching_context(self):
        response = self.client.get('/api/staff/assessment/assignments/')

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('programme_type_code, session, and semester are required', response.data['detail'])

    def test_assignment_list_is_limited_to_the_selected_teaching_context(self):
        response = self.client.get('/api/staff/assessment/assignments/', {
            'programme_type_code': 'ODEL',
            'session': '2025/2026',
            'semester': 'First Semester',
        })

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['count'], 1)
        self.assertEqual(response.data['results'][0]['title'], 'First semester assignment')


class QuizAPITests(APITestCase):
    """Test Quiz REST API endpoints"""
    
    def setUp(self):
        # Create test users
        self.instructor = PortalUser.objects.create(
            external_id='instructor_1',
            full_name='Instructor One',
            is_staff=True
        )
        self.student = PortalUser.objects.create(
            external_id='student_1',
            full_name='Student One',
            roles=['STUDENT'],
        )
        self.session = AcademicSession.objects.create(name='2025/2026')
        self.semester = Semester.objects.create(name='First Semester')
        
        # Create test data
        self.course = CourseCache.objects.create(
            course_external_id=101,
            course_title='Test Course',
            course_code='TEST101'
        )
        
        self.offering = CourseOffering.objects.create(
            course=self.course,
            session=self.session,
            semester=self.semester,
            programme_type_code='ODEL',
        )
        # Setup enrollments
        StaffAssignedCourse.objects.create(
            staff_external_id=self.instructor.external_id,
            course=self.course,
            course_offering=self.offering,
            programme_type_code='ODEL',
            role='instructor'
        )
        self.student_enrollment = StudentRegisteredCourse.objects.create(
            student_external_id=self.student.external_id,
            course=self.course,
            session='2025/2026',
            semester='First Semester'
        )
        sync_patcher = patch(
            'assessment.apis.get_or_sync_student_registered_courses',
            return_value=[self.student_enrollment],
        )
        sync_patcher.start()
        self.addCleanup(sync_patcher.stop)
        
        self.category = QuestionCategory.objects.create(
            course=self.course,
            name='Test Category'
        )
        
        self.question = Question.objects.create(
            category=self.category,
            qtype='shortanswer',
            name='Test Question',
            question_text='What is 2+2?',
            default_mark=Decimal('10.00')
        )
        
        QuestionAnswer.objects.create(
            question=self.question,
            answer_text='4',
            fraction=Decimal('1.0'),
            order=1
        )
        
        self.quiz = Quiz.objects.create(
            course=self.course,
            name='Test Quiz',
            is_published=True,
            time_limit=3600,
            max_grade=Decimal('100.00'),
            max_attempts=3
        )
        
        QuizQuestion.objects.create(
            quiz=self.quiz,
            question=self.question,
            order=1,
            max_mark=Decimal('10.00')
        )
        
        self.client = APIClient()
    
    def test_list_quizzes(self):
        """Test listing quizzes"""
        self.client.force_authenticate(user=self.student)
        response = self.client.get(
            '/api/student/assessment/quizzes/',
            {'session': '2025/2026', 'semester': 'First Semester'},
        )
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data['results']), 1)
        self.assertEqual(response.data['results'][0]['name'], 'Test Quiz')

    def test_staff_quiz_detail_includes_linked_question_slots(self):
        """The staff detail response powers the question-management screen."""
        self.client.force_authenticate(user=self.instructor)

        response = self.client.get(f'/api/staff/assessment/quizzes/{self.quiz.id}/')

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['course_code'], self.course.course_code)
        self.assertEqual(len(response.data['quiz_questions']), 1)
        self.assertEqual(response.data['quiz_questions'][0]['question'], str(self.question.id))

    def test_staff_can_publish_quiz_without_resubmitting_course(self):
        """A partial publication update must not require the create-only course_id."""
        self.quiz.is_published = False
        self.quiz.save(update_fields=['is_published'])
        self.client.force_authenticate(user=self.instructor)

        response = self.client.patch(
            f'/api/staff/assessment/quizzes/{self.quiz.id}/',
            {'is_published': True},
            format='json',
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertTrue(response.data['is_published'])
        self.quiz.refresh_from_db()
        self.assertTrue(self.quiz.is_published)

    def test_staff_can_reorder_every_quiz_question_slot(self):
        second_question = Question.objects.create(
            category=self.category,
            qtype='shortanswer',
            name='Second Question',
            question_text='What is 3+3?',
            default_mark=Decimal('10.00'),
        )
        second_slot = QuizQuestion.objects.create(
            quiz=self.quiz,
            question=second_question,
            order=2,
            max_mark=Decimal('10.00'),
        )
        first_slot = self.quiz.quiz_questions.get(question=self.question)
        self.client.force_authenticate(user=self.instructor)

        response = self.client.post(
            f'/api/staff/assessment/quizzes/{self.quiz.id}/reorder-questions/',
            {'slot_ids': [str(second_slot.id), str(first_slot.id)]},
            format='json',
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual([item['id'] for item in response.data], [str(second_slot.id), str(first_slot.id)])
        self.assertEqual(list(self.quiz.quiz_questions.order_by('order').values_list('id', flat=True)), [second_slot.id, first_slot.id])
    
    def test_create_quiz_instructor_only(self):
        """Test that only instructors can create quizzes"""
        # Student should not be able to create
        self.client.force_authenticate(user=self.student)
        data = {
            'course_id': str(self.course.course_external_id),
            'name': 'New Quiz',
            'max_grade': '100.00'
        }
        response = self.client.post('/api/staff/assessment/quizzes/', data)
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        
        # Instructor should be able to create
        self.client.force_authenticate(user=self.instructor)
        response = self.client.post('/api/staff/assessment/quizzes/', data)
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
    
    def test_start_quiz_attempt(self):
        """Test starting a quiz attempt via API"""
        self.client.force_authenticate(user=self.student)
        
        data = {'session': '2025/2026', 'semester': 'First Semester'}
        response = self.client.post(
            f'/api/student/assessment/quizzes/{self.quiz.id}/start/',
            data,
            format='json'
        )
        
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertIn('id', response.data)
        self.assertEqual(response.data['state'], 'in_progress')

    def test_cannot_start_quiz_without_linked_question_slots(self):
        """A quiz shell must have question-bank slots before an attempt exists."""
        self.quiz.quiz_questions.all().delete()
        self.client.force_authenticate(user=self.student)

        response = self.client.post(
            f'/api/student/assessment/quizzes/{self.quiz.id}/start/',
            {'session': '2025/2026', 'semester': 'First Semester'},
            format='json',
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('no linked questions', response.data['message'])
        self.assertFalse(
            QuizAttempt.objects.filter(quiz=self.quiz, user_external_id=self.student.external_id).exists()
        )

    def test_active_attempt_never_exposes_answer_keys_or_grading(self):
        """Students may restore work, but cannot inspect correctness before submission."""
        self.client.force_authenticate(user=self.student)
        start_response = self.client.post(
            f'/api/student/assessment/quizzes/{self.quiz.id}/start/',
            {'session': '2025/2026', 'semester': 'First Semester'},
            format='json',
        )

        response = self.client.get(f'/api/student/assessment/attempts/{start_response.data["id"]}/')

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        question_attempt = response.data['question_attempts'][0]
        self.assertIsNone(question_attempt['correct_answer'])
        self.assertIsNone(question_attempt['fraction'])
        self.assertIsNone(question_attempt['score'])
        self.assertNotIn('fraction', question_attempt['question']['answers'][0])

    def test_start_closed_quiz_returns_validation_response(self):
        """A closed quiz is a client-visible availability error, not a server error."""
        self.client.force_authenticate(user=self.student)
        self.quiz.time_close = timezone.now() - timedelta(minutes=1)
        self.quiz.save(update_fields=['time_close'])

        response = self.client.post(
            f'/api/student/assessment/quizzes/{self.quiz.id}/start/',
            {'session': '2025/2026', 'semester': 'First Semester'},
            format='json',
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('Quiz closed at', response.data['message'])
        self.assertFalse(
            QuizAttempt.objects.filter(quiz=self.quiz, user_external_id=self.student.external_id).exists()
        )
    
    def test_submit_response_via_api(self):
        """Test submitting a question response via API"""
        self.client.force_authenticate(user=self.student)
        
        # Start attempt
        start_response = self.client.post(
            f'/api/student/assessment/quizzes/{self.quiz.id}/start/',
            {'session': '2025/2026', 'semester': 'First Semester'},
            format='json'
        )
        attempt_id = start_response.data['id']
        
        # Submit response
        data = {
            'question_id': str(self.question.id),
            'response': {'text': '4'}
        }
        response = self.client.post(
            f'/api/student/assessment/quizzes/{self.quiz.id}/attempts/{attempt_id}/submit/',
            data,
            format='json'
        )
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertTrue(response.data['success'])
        self.assertEqual(float(response.data['fraction']), 1.0)
        self.assertEqual(float(response.data['score']), 10.0)
    
    def test_finish_quiz_attempt(self):
        """Test finishing a quiz attempt via API"""
        self.client.force_authenticate(user=self.student)
        
        # Start attempt
        start_response = self.client.post(
            f'/api/student/assessment/quizzes/{self.quiz.id}/start/',
            {'session': '2025/2026', 'semester': 'First Semester'},
            format='json'
        )
        attempt_id = start_response.data['id']
        
        # Submit response
        self.client.post(
            f'/api/student/assessment/quizzes/{self.quiz.id}/attempts/{attempt_id}/submit/',
            {
                'question_id': str(self.question.id),
                'response': {'text': '4'}
            },
            format='json'
        )
        
        # Finish attempt
        response = self.client.post(
            f'/api/student/assessment/quizzes/{self.quiz.id}/attempts/{attempt_id}/finish/',
            format='json'
        )
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['state'], 'finished')
        self.assertIsNotNone(response.data['finished_at'])
    
    def test_list_quiz_attempts(self):
        """Test listing quiz attempts"""
        self.client.force_authenticate(user=self.student)
        
        # Create an attempt
        attempt = QuizService.start_attempt(
            quiz=self.quiz,
            user_external_id=self.student.external_id
        )
        
        response = self.client.get('/api/student/assessment/attempts/')
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data['results']), 1)
        self.assertEqual(response.data['results'][0]['id'], str(attempt.id))

    def test_staff_can_view_paginated_quiz_gradebook(self):
        """Staff gradebook exposes participant details only for assigned-course quizzes."""
        finished_attempt = QuizService.start_attempt(
            quiz=self.quiz,
            user_external_id=self.student.external_id,
        )
        QuizService.submit_response(
            attempt=finished_attempt,
            question=self.question,
            response={'text': '4'},
        )
        QuizService.finish_attempt(finished_attempt)

        self.client.force_authenticate(user=self.instructor)
        response = self.client.get(
            f'/api/staff/assessment/quizzes/{self.quiz.id}/gradebook/',
            {'search': 'Student One', 'page_size': 10},
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['count'], 1)
        self.assertEqual(response.data['summary']['completed_count'], 1)
        self.assertEqual(response.data['results'][0]['participant_name'], 'Student One')
        self.assertEqual(response.data['results'][0]['participant_email'], None)
        self.assertEqual(float(response.data['results'][0]['max_grade']), 100.0)
    
    def test_manual_grading(self):
        """Test manual grading via API (instructor only)"""
        # Create essay question
        essay_question = Question.objects.create(
            category=self.category,
            qtype='essay',
            name='Essay Question',
            question_text='Explain Python',
            default_mark=Decimal('20.00')
        )
        
        QuizQuestion.objects.create(
            quiz=self.quiz,
            question=essay_question,
            order=2,
            max_mark=Decimal('20.00')
        )
        
        # Student starts attempt and submits essay
        attempt = QuizService.start_attempt(
            quiz=self.quiz,
            user_external_id=self.student.external_id
        )
        
        question_attempt = QuizService.submit_response(
            attempt=attempt,
            question=essay_question,
            response={'text': 'Python is great because...'}
        )
        
        # Instructor grades the essay
        self.client.force_authenticate(user=self.instructor)
        data = {
            'fraction': '0.85',
            'feedback': 'Good work, but missing some points'
        }
        response = self.client.post(
            f'/api/staff/assessment/attempts/{attempt.id}/questions/{question_attempt.id}/grade/',
            data,
            format='json'
        )
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertTrue(response.data['success'])
        self.assertEqual(float(response.data['fraction']), 0.85)
    
    def test_student_cannot_manual_grade(self):
        """Test that students cannot manually grade"""
        attempt = QuizService.start_attempt(
            quiz=self.quiz,
            user_external_id=self.student.external_id
        )
        
        question_attempt = attempt.question_attempts.first()
        
        self.client.force_authenticate(user=self.student)
        data = {'fraction': '1.0', 'feedback': 'Perfect'}
        response = self.client.post(
            f'/api/staff/assessment/attempts/{attempt.id}/questions/{question_attempt.id}/grade/',
            data,
            format='json'
        )
        
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_student_can_complete_an_auto_graded_assessment_from_the_separate_bank(self):
        assessment_category = QuestionCategory.objects.create(
            course=self.course,
            bank='assessment',
            name='Assessment Fundamentals',
        )
        question = Question.objects.create(
            category=assessment_category,
            qtype='multichoice',
            name='Binary MCQ',
            question_text='Which value represents true in a Boolean expression?',
            default_mark=Decimal('5.00'),
        )
        correct_answer = QuestionAnswer.objects.create(
            question=question,
            answer_text='True',
            fraction=Decimal('1.00'),
            order=1,
        )
        QuestionAnswer.objects.create(
            question=question,
            answer_text='False',
            fraction=Decimal('0.00'),
            order=2,
        )
        assessment = Assessment.objects.create(
            course=self.course,
            course_offering=self.offering,
            name='Boolean checkpoint',
            is_published=True,
            max_grade=Decimal('20.00'),
            show_feedback=True,
        )
        AssessmentQuestion.objects.create(
            assessment=assessment,
            question=question,
            order=1,
            max_mark=Decimal('5.00'),
        )

        self.client.force_authenticate(user=self.student)
        period = {'session': '2025/2026', 'semester': 'First Semester'}
        start_response = self.client.post(
            f'/api/student/assessment/assessments/{assessment.id}/start/',
            period,
            format='json',
        )

        self.assertEqual(start_response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(start_response.data['question_attempts'][0]['question']['id'], str(question.id))
        self.assertIsNone(start_response.data['question_attempts'][0]['fraction'])
        attempt_id = start_response.data['id']

        submit_response = self.client.post(
            f'/api/student/assessment/assessments/{assessment.id}/attempts/{attempt_id}/submit/',
            {
                'question_id': str(question.id),
                'response': {'selected': [str(correct_answer.id)]},
            },
            format='json',
        )
        self.assertEqual(submit_response.status_code, status.HTTP_200_OK)

        finish_response = self.client.post(
            f'/api/student/assessment/assessments/{assessment.id}/attempts/{attempt_id}/finish/',
            format='json',
        )
        self.assertEqual(finish_response.status_code, status.HTTP_200_OK)
        self.assertEqual(finish_response.data['state'], 'finished')
        self.assertEqual(finish_response.data['total_score'], 20.0)
        self.assertTrue(Grade.objects.filter(
            assessment_attempt_id=attempt_id,
            grade_type='assessment',
        ).exists())

    def test_student_can_leave_and_resume_a_self_paced_assessment(self):
        """Starting again restores the saved draft and does not use an attempt."""
        category = QuestionCategory.objects.create(
            course=self.course, bank='assessment', name='Resume assessment bank',
        )
        question = Question.objects.create(
            category=category, qtype='truefalse', name='Resume question',
            question_text='Saved work remains available.', default_mark=Decimal('1.00'),
        )
        correct = QuestionAnswer.objects.create(
            question=question, answer_text='True', fraction=Decimal('1.00'), order=1,
        )
        QuestionAnswer.objects.create(
            question=question, answer_text='False', fraction=Decimal('0.00'), order=2,
        )
        assessment = Assessment.objects.create(
            course=self.course,
            course_offering=self.offering,
            name='Self-paced checkpoint',
            is_published=True,
            max_grade=Decimal('10.00'),
            max_attempts=1,
            # This legacy configuration must not give assessment attempts a
            # deadline now that they are self-paced.
            time_limit=60,
        )
        AssessmentQuestion.objects.create(
            assessment=assessment, question=question, order=1, max_mark=Decimal('1.00'),
        )

        self.client.force_authenticate(user=self.student)
        period = {'session': '2025/2026', 'semester': 'First Semester'}
        started = self.client.post(
            f'/api/student/assessment/assessments/{assessment.id}/start/', period, format='json',
        )
        self.assertEqual(started.status_code, status.HTTP_201_CREATED)
        self.assertIsNone(started.data['deadline_at'])

        saved = self.client.post(
            f'/api/student/assessment/assessments/{assessment.id}/attempts/{started.data["id"]}/submit/',
            {'question_id': str(question.id), 'response': {'selected': str(correct.id)}}, format='json',
        )
        self.assertEqual(saved.status_code, status.HTTP_200_OK)

        added_question = Question.objects.create(
            category=category, qtype='truefalse', name='Later question',
            question_text='A lecturer may add this after the student starts.', default_mark=Decimal('1.00'),
        )
        QuestionAnswer.objects.create(
            question=added_question, answer_text='True', fraction=Decimal('1.00'), order=1,
        )
        QuestionAnswer.objects.create(
            question=added_question, answer_text='False', fraction=Decimal('0.00'), order=2,
        )
        AssessmentQuestion.objects.create(
            assessment=assessment, question=added_question, order=2, max_mark=Decimal('1.00'),
        )

        resumed = self.client.post(
            f'/api/student/assessment/assessments/{assessment.id}/start/', period, format='json',
        )
        self.assertEqual(resumed.status_code, status.HTTP_201_CREATED)
        self.assertEqual(resumed.data['id'], started.data['id'])
        self.assertEqual(resumed.data['state'], 'in_progress')
        self.assertEqual(len(resumed.data['question_attempts']), 2)
        self.assertEqual(resumed.data['question_attempts'][0]['response']['selected'], str(correct.id))
        self.assertEqual(resumed.data['question_attempts'][1]['question']['id'], str(added_question.id))
        self.assertEqual(
            AssessmentAttempt.objects.filter(
                assessment=assessment, user_external_id=self.student.external_id,
            ).count(),
            1,
        )

    def test_assessment_attempt_uses_frozen_question_and_mark_snapshot(self):
        """Editing a bank item later must never change an in-progress result."""
        category = QuestionCategory.objects.create(
            course=self.course, bank='assessment', name='Snapshot bank',
        )
        question = Question.objects.create(
            category=category, qtype='truefalse', name='Snapshot question',
            question_text='The original answer remains true.', default_mark=Decimal('2.00'),
        )
        correct = QuestionAnswer.objects.create(
            question=question, answer_text='True', fraction=Decimal('1.00'), order=1,
        )
        QuestionAnswer.objects.create(
            question=question, answer_text='False', fraction=Decimal('0.00'), order=2,
        )
        assessment = Assessment.objects.create(
            course=self.course, course_offering=self.offering, name='Snapshot checkpoint',
            is_published=True, max_grade=Decimal('10.00'), show_feedback=True,
        )
        AssessmentQuestion.objects.create(
            assessment=assessment, question=question, order=1, max_mark=Decimal('2.00'),
        )
        self.client.force_authenticate(user=self.student)
        period = {'session': '2025/2026', 'semester': 'First Semester'}
        started = self.client.post(
            f'/api/student/assessment/assessments/{assessment.id}/start/', period, format='json',
        )
        self.assertEqual(started.status_code, status.HTTP_201_CREATED)
        question.answers.filter(id=correct.id).update(fraction=Decimal('0.00'))
        AssessmentQuestion.objects.filter(assessment=assessment, question=question).update(max_mark=Decimal('99.00'))
        attempt_id = started.data['id']
        saved = self.client.post(
            f'/api/student/assessment/assessments/{assessment.id}/attempts/{attempt_id}/submit/',
            {'question_id': str(question.id), 'response': {'selected': str(correct.id)}}, format='json',
        )
        self.assertEqual(saved.status_code, status.HTTP_200_OK)
        finished = self.client.post(
            f'/api/student/assessment/assessments/{assessment.id}/attempts/{attempt_id}/finish/', format='json',
        )
        self.assertEqual(finished.status_code, status.HTTP_200_OK)
        self.assertEqual(finished.data['total_score'], 10.0)
        # The retry is safe even when the browser's deadline and submit button race.
        finished = self.client.post(
            f'/api/student/assessment/assessments/{assessment.id}/attempts/{attempt_id}/finish/', format='json',
        )
        self.assertEqual(finished.status_code, status.HTTP_200_OK)
        self.assertTrue(finished.data['show_feedback'])
        self.assertIsNone(finished.data['total_score'])

    def test_teacher_controls_score_release_for_finished_assessments(self):
        """Scores stay hidden until the lecturer releases them, even after submission."""
        category = QuestionCategory.objects.create(
            course=self.course, bank='assessment', name='Feedback snapshot bank',
        )
        question = Question.objects.create(
            category=category, qtype='truefalse', name='Feedback snapshot question',
            question_text='This score is frozen.', default_mark=Decimal('1.00'),
        )
        correct = QuestionAnswer.objects.create(
            question=question, answer_text='True', fraction=Decimal('1.00'), order=1,
        )
        QuestionAnswer.objects.create(
            question=question, answer_text='False', fraction=Decimal('0.00'), order=2,
        )
        assessment = Assessment.objects.create(
            course=self.course, course_offering=self.offering, name='Feedback snapshot',
            is_published=True, max_grade=Decimal('10.00'), show_feedback=True,
        )
        AssessmentQuestion.objects.create(
            assessment=assessment, question=question, order=1, max_mark=Decimal('1.00'),
        )
        self.client.force_authenticate(user=self.student)
        started = self.client.post(
            f'/api/student/assessment/assessments/{assessment.id}/start/',
            {'session': '2025/2026', 'semester': 'First Semester'}, format='json',
        )
        self.assertEqual(started.status_code, status.HTTP_201_CREATED)
        assessment.show_feedback = False
        assessment.save(update_fields=['show_feedback'])
        attempt_id = started.data['id']
        self.client.post(
            f'/api/student/assessment/assessments/{assessment.id}/attempts/{attempt_id}/submit/',
            {'question_id': str(question.id), 'response': {'selected': str(correct.id)}}, format='json',
        )
        completed = self.client.post(
            f'/api/student/assessment/assessments/{assessment.id}/attempts/{attempt_id}/finish/', format='json',
        )
        self.assertEqual(completed.status_code, status.HTTP_200_OK)
        self.assertFalse(completed.data['show_feedback'])
        self.assertIsNone(completed.data['total_score'])

        assessment.show_feedback = True
        assessment.save(update_fields=['show_feedback'])
        released = self.client.get(
            f'/api/student/assessment/assessment-attempts/{attempt_id}/',
        )
        self.assertEqual(released.status_code, status.HTTP_200_OK)
        self.assertTrue(released.data['show_feedback'])
        self.assertEqual(released.data['total_score'], 10.0)

        self.client.force_authenticate(user=self.instructor)
        staff_activity = self.client.get(
            f'/api/staff/assessment/assessment-attempts/{attempt_id}/',
            {'programme_type_code': 'ODEL', 'session': '2025/2026', 'semester': 'First Semester'},
        )
        self.assertEqual(staff_activity.status_code, status.HTTP_200_OK)
        self.assertEqual(staff_activity.data['user_external_id'], self.student.external_id)
        self.assertEqual(staff_activity.data['question_attempts'][0]['response']['selected'], str(correct.id))
        self.assertEqual(staff_activity.data['question_attempts'][0]['question']['id'], str(question.id))

    def test_assessment_essay_is_saved_and_can_be_manually_marked(self):
        """Assessment essays are retained, then graded by the assigned lecturer."""
        category = QuestionCategory.objects.create(
            course=self.course, bank='assessment', name='Essay assessment bank',
        )
        essay = Question.objects.create(
            category=category, qtype='essay', name='Explain normalization',
            question_text='Explain why database normalization matters.', default_mark=Decimal('10.00'),
        )
        assessment = Assessment.objects.create(
            course=self.course, course_offering=self.offering, name='Written checkpoint',
            is_published=True, max_grade=Decimal('10.00'), show_feedback=True,
        )
        AssessmentQuestion.objects.create(
            assessment=assessment, question=essay, order=1, max_mark=Decimal('10.00'),
        )

        self.client.force_authenticate(user=self.student)
        period = {'session': '2025/2026', 'semester': 'First Semester'}
        started = self.client.post(
            f'/api/student/assessment/assessments/{assessment.id}/start/', period, format='json',
        )
        self.assertEqual(started.status_code, status.HTTP_201_CREATED)
        attempt_id = started.data['id']
        saved = self.client.post(
            f'/api/student/assessment/assessments/{assessment.id}/attempts/{attempt_id}/submit/',
            {'question_id': str(essay.id), 'response': {'text': 'It reduces duplication and update anomalies.'}},
            format='json',
        )
        self.assertEqual(saved.status_code, status.HTTP_200_OK)
        self.assertEqual(
            self.client.post(
                f'/api/student/assessment/assessments/{assessment.id}/attempts/{attempt_id}/finish/', format='json',
            ).status_code,
            status.HTTP_200_OK,
        )

        question_attempt = AssessmentQuestionAttempt.objects.get(assessment_attempt_id=attempt_id, question=essay)
        self.assertIsNone(question_attempt.score)
        self.assertIsNone(AssessmentAttempt.objects.get(id=attempt_id).total_score)
        self.assertFalse(Grade.objects.filter(assessment_attempt_id=attempt_id).exists())
        self.client.force_authenticate(user=self.instructor)
        marked = self.client.post(
            f'/api/staff/assessment/assessment-attempts/{attempt_id}/questions/{question_attempt.id}/grade/',
            {'fraction': '0.80', 'feedback': 'Strong explanation; add an example next time.'},
            format='json',
        )
        self.assertEqual(marked.status_code, status.HTTP_200_OK)
        self.assertEqual(marked.data['score'], 8.0)
        question_attempt.refresh_from_db()
        self.assertTrue(question_attempt.manually_graded)
        self.assertEqual(question_attempt.score, Decimal('8.00'))
        self.assertEqual(question_attempt.feedback, 'Strong explanation; add an example next time.')
        self.assertEqual(AssessmentAttempt.objects.get(id=attempt_id).total_score, Decimal('8.00'))
        self.assertTrue(Grade.objects.filter(assessment_attempt_id=attempt_id, marks=Decimal('8.00')).exists())

    def test_assessment_manual_grading_requires_the_assigned_offering(self):
        """A same-course assignment in another term cannot mark this attempt."""
        category = QuestionCategory.objects.create(
            course=self.course, bank='assessment', name='Offering-scoped essay bank',
        )
        essay = Question.objects.create(
            category=category, qtype='essay', name='Term-specific essay',
            question_text='Explain the first-term topic.', default_mark=Decimal('10.00'),
        )
        assessment = Assessment.objects.create(
            course=self.course, course_offering=self.offering, name='First-term essay', is_published=True,
        )
        AssessmentQuestion.objects.create(
            assessment=assessment, question=essay, order=1, max_mark=Decimal('10.00'),
        )
        attempt = AssessmentService.start_attempt(assessment=assessment, user_external_id=self.student.external_id)
        AssessmentService.submit_response(attempt=attempt, question=essay, response={'text': 'My response.'})
        AssessmentService.finish_attempt(attempt)
        question_attempt = AssessmentQuestionAttempt.objects.get(assessment_attempt=attempt, question=essay)

        second_semester = Semester.objects.create(name='Second Semester')
        second_offering = CourseOffering.objects.create(
            course=self.course, session=self.session, semester=second_semester, programme_type_code='ODEL',
        )
        StaffAssignedCourse.objects.filter(
            staff_external_id=self.instructor.external_id,
            course=self.course,
        ).delete()
        StaffAssignedCourse.objects.create(
            staff_external_id=self.instructor.external_id,
            course=self.course,
            course_offering=second_offering,
            programme_type_code='ODEL',
            role='INSTRUCTOR',
        )
        self.client.force_authenticate(user=self.instructor)
        response = self.client.post(
            f'/api/staff/assessment/assessment-attempts/{attempt.id}/questions/{question_attempt.id}/grade/',
            {'fraction': '1.00'},
            format='json',
        )

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        question_attempt.refresh_from_db()
        self.assertFalse(question_attempt.manually_graded)

    def test_assessment_score_waits_for_every_essay_mark(self):
        """One marked essay must not publish a partial final assessment score."""
        category = QuestionCategory.objects.create(
            course=self.course, bank='assessment', name='Multiple essay assessment bank',
        )
        first_essay = Question.objects.create(
            category=category, qtype='essay', name='First essay', question_text='First response.', default_mark=Decimal('5.00'),
        )
        second_essay = Question.objects.create(
            category=category, qtype='essay', name='Second essay', question_text='Second response.', default_mark=Decimal('5.00'),
        )
        assessment = Assessment.objects.create(
            course=self.course, course_offering=self.offering, name='Two essays', max_grade=Decimal('10.00'),
        )
        AssessmentQuestion.objects.create(assessment=assessment, question=first_essay, order=1, max_mark=Decimal('5.00'))
        AssessmentQuestion.objects.create(assessment=assessment, question=second_essay, order=2, max_mark=Decimal('5.00'))
        attempt = AssessmentService.start_attempt(assessment=assessment, user_external_id=self.student.external_id)
        AssessmentService.submit_response(attempt=attempt, question=first_essay, response={'text': 'First answer.'})
        AssessmentService.submit_response(attempt=attempt, question=second_essay, response={'text': 'Second answer.'})
        AssessmentService.finish_attempt(attempt)

        first_attempt = AssessmentQuestionAttempt.objects.get(assessment_attempt=attempt, question=first_essay)
        second_attempt = AssessmentQuestionAttempt.objects.get(assessment_attempt=attempt, question=second_essay)
        AssessmentService.manually_grade_question_attempt(first_attempt, Decimal('1.00'))
        attempt.refresh_from_db()
        self.assertIsNone(attempt.total_score)
        self.assertFalse(Grade.objects.filter(assessment_attempt=attempt).exists())

        AssessmentService.manually_grade_question_attempt(second_attempt, Decimal('0.50'))
        attempt.refresh_from_db()
        self.assertEqual(attempt.total_score, Decimal('7.50'))
        self.assertTrue(Grade.objects.filter(assessment_attempt=attempt, marks=Decimal('7.50')).exists())

    def test_assessment_rejects_non_positive_slot_marks_through_the_api(self):
        category = QuestionCategory.objects.create(
            course=self.course, bank='assessment', name='Slot validation bank',
        )
        question = Question.objects.create(
            category=category, qtype='shortanswer', name='Slot question',
            question_text='Name a data type.', default_mark=Decimal('1.00'),
        )
        QuestionAnswer.objects.create(question=question, answer_text='String', fraction=Decimal('1.00'), order=1)
        assessment = Assessment.objects.create(course=self.course, course_offering=self.offering, name='Slot validation')
        self.client.force_authenticate(user=self.instructor)
        response = self.client.post(
            '/api/staff/assessment/assessments/questions/?programme_type_code=ODEL&session=2025/2026&semester=First%20Semester',
            {'assessment': str(assessment.id), 'question': str(question.id), 'order': 1, 'max_mark': '0'},
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('max_mark', response.data)

    def test_assessment_allows_appending_questions_but_locks_existing_slots_after_start(self):
        """Started attempts keep their snapshot while staff can prepare later attempts."""
        category = QuestionCategory.objects.create(
            course=self.course, bank='assessment', name='Append-only bank',
        )
        first_question = Question.objects.create(
            category=category, qtype='truefalse', name='First append-only question',
            question_text='The first question is retained.', default_mark=Decimal('1.00'),
        )
        QuestionAnswer.objects.create(
            question=first_question, answer_text='True', fraction=Decimal('1.00'), order=1,
        )
        QuestionAnswer.objects.create(
            question=first_question, answer_text='False', fraction=Decimal('0.00'), order=2,
        )
        assessment = Assessment.objects.create(
            course=self.course, course_offering=self.offering, name='Append-only checkpoint',
            is_published=True, max_grade=Decimal('10.00'),
        )
        first_slot = AssessmentQuestion.objects.create(
            assessment=assessment, question=first_question, order=1, max_mark=Decimal('1.00'),
        )

        self.client.force_authenticate(user=self.student)
        started = self.client.post(
            f'/api/student/assessment/assessments/{assessment.id}/start/',
            {'session': '2025/2026', 'semester': 'First Semester'}, format='json',
        )
        self.assertEqual(started.status_code, status.HTTP_201_CREATED)
        attempt = AssessmentAttempt.objects.get(id=started.data['id'])
        self.assertEqual(attempt.question_attempts.count(), 1)

        later_question = Question.objects.create(
            category=category, qtype='truefalse', name='Later append-only question',
            question_text='This is added for later attempts.', default_mark=Decimal('1.00'),
        )
        QuestionAnswer.objects.create(
            question=later_question, answer_text='True', fraction=Decimal('1.00'), order=1,
        )
        QuestionAnswer.objects.create(
            question=later_question, answer_text='False', fraction=Decimal('0.00'), order=2,
        )
        staff_period = 'programme_type_code=ODEL&session=2025/2026&semester=First%20Semester'
        self.client.force_authenticate(user=self.instructor)
        added = self.client.post(
            f'/api/staff/assessment/assessments/questions/?{staff_period}',
            {
                'assessment': str(assessment.id),
                'question': str(later_question.id),
                'order': 2,
                'max_mark': '1.00',
            },
            format='json',
        )
        self.assertEqual(added.status_code, status.HTTP_201_CREATED)
        self.assertEqual(assessment.assessment_questions.count(), 2)
        self.assertEqual(attempt.question_attempts.count(), 1)

        removed = self.client.delete(
            f'/api/staff/assessment/assessments/questions/{first_slot.id}/?{staff_period}',
        )
        self.assertEqual(removed.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertTrue(AssessmentQuestion.objects.filter(id=first_slot.id).exists())

    def test_assessment_detail_is_limited_to_the_selected_staff_offering(self):
        second_semester = Semester.objects.create(name='Second Semester')
        second_offering = CourseOffering.objects.create(
            course=self.course, session=self.session, semester=second_semester, programme_type_code='ODEL',
        )
        StaffAssignedCourse.objects.create(
            staff_external_id=self.instructor.external_id, course=self.course,
            course_offering=second_offering, programme_type_code='ODEL', role='INSTRUCTOR',
        )
        assessment = Assessment.objects.create(course=self.course, course_offering=self.offering, name='First-term only')
        self.client.force_authenticate(user=self.instructor)
        response = self.client.get(
            f'/api/staff/assessment/assessments/{assessment.id}/',
            {'programme_type_code': 'ODEL', 'session': '2025/2026', 'semester': 'Second Semester'},
        )
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    def test_student_assessment_list_excludes_ambiguous_legacy_enrolment(self):
        """A course shared by two programme offerings must not leak either assessment."""
        CourseOffering.objects.create(
            course=self.course, session=self.session, semester=self.semester, programme_type_code='CAMPUS',
        )
        Assessment.objects.create(
            course=self.course, course_offering=self.offering, name='ODEL-only checkpoint', is_published=True,
        )
        self.client.force_authenticate(user=self.student)
        response = self.client.get(
            '/api/student/assessment/assessments/',
            {'session': '2025/2026', 'semester': 'First Semester'},
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['count'], 0)

    def test_assessment_with_slots_cannot_be_moved_to_another_course(self):
        other_course = CourseCache.objects.create(course_external_id=909, course_title='Other course', course_code='OTHER909')
        other_offering = CourseOffering.objects.create(
            course=other_course, session=self.session, semester=self.semester, programme_type_code='ODEL',
        )
        StaffAssignedCourse.objects.create(
            staff_external_id=self.instructor.external_id, course=other_course,
            course_offering=other_offering, programme_type_code='ODEL', role='INSTRUCTOR',
        )
        category = QuestionCategory.objects.create(course=self.course, bank='assessment', name='Move protection bank')
        question = Question.objects.create(category=category, qtype='shortanswer', name='Move protection question', question_text='Answer.', default_mark=Decimal('1.00'))
        QuestionAnswer.objects.create(question=question, answer_text='Answer', fraction=Decimal('1.00'), order=1)
        assessment = Assessment.objects.create(course=self.course, course_offering=self.offering, name='Cannot move')
        AssessmentQuestion.objects.create(assessment=assessment, question=question, order=1, max_mark=Decimal('1.00'))
        self.client.force_authenticate(user=self.instructor)
        response = self.client.patch(
            f'/api/staff/assessment/assessments/{assessment.id}/?programme_type_code=ODEL&session=2025/2026&semester=First%20Semester',
            {'course_id': str(other_course.course_external_id)}, format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('course_id', response.data)

    def test_assessment_gradebook_reports_assessment_breakdown(self):
        category = QuestionCategory.objects.create(course=self.course, bank='assessment', name='Gradebook assessment bank')
        question = Question.objects.create(category=category, qtype='truefalse', name='Gradebook question', question_text='A is A.', default_mark=Decimal('1.00'))
        correct = QuestionAnswer.objects.create(question=question, answer_text='True', fraction=Decimal('1.00'), order=1)
        QuestionAnswer.objects.create(question=question, answer_text='False', fraction=Decimal('0.00'), order=2)
        assessment = Assessment.objects.create(course=self.course, course_offering=self.offering, name='Gradebook assessment', max_grade=Decimal('20.00'))
        AssessmentQuestion.objects.create(assessment=assessment, question=question, order=1, max_mark=Decimal('1.00'))
        attempt = AssessmentService.start_attempt(assessment=assessment, user_external_id=self.student.external_id)
        AssessmentService.submit_response(attempt=attempt, question=question, response={'selected': str(correct.id)})
        AssessmentService.finish_attempt(attempt)
        totals = calculate_student_total(self.student.external_id, self.course.id)
        summary = get_gradebook_summary(self.course.id)[0]
        self.assertEqual(totals['assessment_marks'], Decimal('20.00'))
        self.assertEqual(totals['assessment_possible'], Decimal('20.00'))
        self.assertEqual(summary['assessment_count'], 1)

    def test_assessment_rejects_quiz_question_bank_items(self):
        assessment = Assessment.objects.create(course=self.course, name='Separate bank check')
        with self.assertRaises(ValidationError):
            AssessmentQuestion.objects.create(
                assessment=assessment,
                question=self.question,
                order=1,
                max_mark=Decimal('10.00'),
            )


class QuestionAPITests(APITestCase):
    """Test Question Bank API endpoints"""
    
    def setUp(self):
        self.instructor = PortalUser.objects.create(
            external_id='instructor_1',
            full_name='Instructor One',
            is_staff=True
        )
        self.student = PortalUser.objects.create(
            external_id='student_1',
            full_name='Student One'
        )
        
        self.course = CourseCache.objects.create(
            course_external_id=101,
            course_title='Test Course',
            course_code='TEST101'
        )
        
        # Setup enrollments
        StaffAssignedCourse.objects.create(
            staff_external_id=self.instructor.external_id,
            course=self.course,
            role='instructor'
        )
        
        self.category = QuestionCategory.objects.create(
            course=self.course,
            name='Test Category'
        )
        
        self.client = APIClient()
    
    def test_create_question_with_answers(self):
        """Test creating a question with nested answers"""
        self.client.force_authenticate(user=self.instructor)
        
        data = {
            'category': str(self.category.id),
            'qtype': 'multichoice',
            'name': 'Python Variables',
            'question_text': 'Which are valid Python variables?',
            'default_mark': '10.00',
            'answers': [
                {'answer_text': 'my_var', 'fraction': '0.5', 'order': 1},
                {'answer_text': '_private', 'fraction': '0.5', 'order': 2},
                {'answer_text': '2nd_var', 'fraction': '0.0', 'order': 3}
            ]
        }
        
        response = self.client.post(
            '/api/staff/assessment/question-bank/questions/',
            data,
            format='json'
        )
        
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        
        # Verify answers were created
        question_id = response.data['id']
        question = Question.objects.get(id=question_id)
        self.assertEqual(question.answers.count(), 3)
    
    def test_student_cannot_create_question(self):
        """Test that students cannot create questions"""
        self.client.force_authenticate(user=self.student)
        
        data = {
            'category': str(self.category.id),
            'qtype': 'shortanswer',
            'name': 'Test',
            'question_text': 'Test?',
            'default_mark': '5.00'
        }
        
        response = self.client.post(
            '/api/staff/assessment/question-bank/questions/',
            data,
            format='json'
        )
        
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
    
    def test_list_questions_filtered_by_category(self):
        """Test filtering questions by category"""
        self.client.force_authenticate(user=self.instructor)
        
        # Create questions in different categories
        category2 = QuestionCategory.objects.create(
            course=self.course,
            name='Category 2'
        )
        
        Question.objects.create(
            category=self.category,
            qtype='shortanswer',
            name='Q1',
            question_text='Test 1',
            default_mark=Decimal('5.00')
        )
        
        Question.objects.create(
            category=category2,
            qtype='shortanswer',
            name='Q2',
            question_text='Test 2',
            default_mark=Decimal('5.00')
        )
        
        # Filter by category
        response = self.client.get(
            f'/api/staff/assessment/question-bank/questions/?category={self.category.id}'
        )
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data['results']), 1)
        self.assertEqual(response.data['results'][0]['name'], 'Q1')

    def test_question_bank_filters_categories_and_questions_by_course(self):
        other_course = CourseCache.objects.create(
            course_external_id=202,
            course_title='Other Test Course',
            course_code='TEST202',
        )
        StaffAssignedCourse.objects.create(
            staff_external_id=self.instructor.external_id,
            course=other_course,
            role='instructor',
        )
        other_category = QuestionCategory.objects.create(
            course=other_course,
            name='Other Course Category',
        )
        Question.objects.create(
            category=self.category,
            qtype='shortanswer',
            name='Current Course Question',
            question_text='Question for the current course',
            default_mark=Decimal('5.00'),
        )
        Question.objects.create(
            category=other_category,
            qtype='shortanswer',
            name='Other Course Question',
            question_text='Question for another course',
            default_mark=Decimal('5.00'),
        )

        self.client.force_authenticate(user=self.instructor)
        categories_response = self.client.get(
            f'/api/staff/assessment/question-bank/categories/?course={self.course.course_external_id}'
        )
        questions_response = self.client.get(
            f'/api/staff/assessment/question-bank/questions/?course={self.course.course_external_id}'
        )

        self.assertEqual(categories_response.status_code, status.HTTP_200_OK)
        self.assertEqual(categories_response.data['count'], 1)
        self.assertEqual(categories_response.data['results'][0]['name'], self.category.name)
        self.assertEqual(questions_response.status_code, status.HTTP_200_OK)
        self.assertEqual(questions_response.data['count'], 1)
        self.assertEqual(questions_response.data['results'][0]['name'], 'Current Course Question')

    def test_staff_can_copy_same_course_assessment_questions_to_quiz_category(self):
        source_category = QuestionCategory.objects.create(
            course=self.course,
            bank='assessment',
            name='Assessment source',
        )
        source_question = Question.objects.create(
            category=source_category,
            qtype='singlechoice',
            name='Copied question',
            question_text='Which option is correct?',
            general_feedback='Review the topic notes.',
            default_mark=Decimal('2.00'),
            penalty=Decimal('0.33'),
        )
        QuestionAnswer.objects.create(
            question=source_question,
            answer_text='Correct option',
            fraction=Decimal('1.00'),
            feedback='Exactly right.',
            order=1,
        )
        QuestionAnswer.objects.create(
            question=source_question,
            answer_text='Incorrect option',
            fraction=Decimal('0.00'),
            feedback='Try again.',
            order=2,
        )

        self.client.force_authenticate(user=self.instructor)
        response = self.client.post(
            f'/api/staff/assessment/question-bank/categories/{self.category.id}/copy-from-assessment/',
            {'question_ids': [str(source_question.id)]},
            format='json',
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(len(response.data), 1)
        copied_question = Question.objects.get(id=response.data[0]['id'])
        self.assertNotEqual(copied_question.id, source_question.id)
        self.assertEqual(copied_question.category, self.category)
        self.assertEqual(copied_question.qtype, source_question.qtype)
        self.assertEqual(copied_question.question_text, source_question.question_text)
        self.assertEqual(copied_question.answers.count(), 2)
        self.assertEqual(copied_question.answers.first().answer_text, 'Correct option')
        self.assertEqual(source_question.category, source_category)

    def test_copy_rejects_assessment_question_from_another_course(self):
        other_course = CourseCache.objects.create(
            course_external_id=202,
            course_title='Other Test Course',
            course_code='TEST202',
        )
        StaffAssignedCourse.objects.create(
            staff_external_id=self.instructor.external_id,
            course=other_course,
            role='instructor',
        )
        other_category = QuestionCategory.objects.create(
            course=other_course,
            bank='assessment',
            name='Other assessment source',
        )
        other_question = Question.objects.create(
            category=other_category,
            qtype='shortanswer',
            name='Other course question',
            question_text='Not available for this target.',
            default_mark=Decimal('1.00'),
        )

        self.client.force_authenticate(user=self.instructor)
        response = self.client.post(
            f'/api/staff/assessment/question-bank/categories/{self.category.id}/copy-from-assessment/',
            {'question_ids': [str(other_question.id)]},
            format='json',
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(Question.objects.filter(category=self.category).count(), 0)


class IntegrationTests(TestCase):
    """End-to-end integration tests"""
    
    def test_complete_quiz_workflow(self):
        """Test complete quiz workflow from creation to completion"""
        # Setup
        course = CourseCache.objects.create(
            course_external_id=101,
            course_title='Test Course',
            course_code='TEST101'
        )
        
        category = QuestionCategory.objects.create(
            course=course,
            name='Math'
        )
        
        # Create multiple questions
        q1 = Question.objects.create(
            category=category,
            qtype='shortanswer',
            name='Addition',
            question_text='What is 2+2?',
            default_mark=Decimal('10.00')
        )
        QuestionAnswer.objects.create(
            question=q1,
            answer_text='4',
            fraction=Decimal('1.0'),
            order=1
        )
        
        q2 = Question.objects.create(
            category=category,
            qtype='multichoice',
            name='Even Numbers',
            question_text='Which are even?',
            default_mark=Decimal('10.00')
        )
        ans1 = QuestionAnswer.objects.create(
            question=q2,
            answer_text='2',
            fraction=Decimal('0.5'),
            order=1
        )
        ans2 = QuestionAnswer.objects.create(
            question=q2,
            answer_text='4',
            fraction=Decimal('0.5'),
            order=2
        )
        QuestionAnswer.objects.create(
            question=q2,
            answer_text='3',
            fraction=Decimal('0.0'),
            order=3
        )
        
        # Create quiz
        quiz = Quiz.objects.create(
            course=course,
            name='Math Quiz',
            max_grade=Decimal('100.00'),
            max_attempts=2
        )
        
        QuizQuestion.objects.create(
            quiz=quiz,
            question=q1,
            order=1,
            max_mark=Decimal('10.00')
        )
        
        QuizQuestion.objects.create(
            quiz=quiz,
            question=q2,
            order=2,
            max_mark=Decimal('10.00')
        )
        
        # Student starts quiz
        attempt = QuizService.start_attempt(
            quiz=quiz,
            user_external_id='student_1'
        )
        
        self.assertEqual(attempt.state, 'in_progress')
        self.assertEqual(attempt.question_attempts.count(), 2)
        
        # Student answers question 1 correctly
        qa1 = QuizService.submit_response(
            attempt=attempt,
            question=q1,
            response={'text': '4'}
        )
        self.assertEqual(qa1.fraction, Decimal('1.0'))
        self.assertEqual(qa1.score, Decimal('10.00'))
        
        # Student answers question 2 partially (only one correct)
        qa2 = QuizService.submit_response(
            attempt=attempt,
            question=q2,
            response={'selected': [str(ans1.id)]}
        )
        self.assertEqual(qa2.fraction, Decimal('0.5'))
        self.assertEqual(qa2.score, Decimal('5.00'))
        
        # Student finishes quiz
        QuizService.finish_attempt(attempt)
        
        attempt.refresh_from_db()
        self.assertEqual(attempt.state, 'finished')
        # Score: 15 out of 20 possible = 75%, scaled to max_grade of 100
        self.assertEqual(attempt.total_score, Decimal('75.00'))
        
        # Verify summary
        summary = QuizService.get_attempt_summary(attempt)
        self.assertEqual(summary['total_questions'], 2)
        self.assertEqual(summary['answered_questions'], 2)
        # Score is scaled to max_grade: 15/20 = 75%
        self.assertEqual(summary['total_score'], 75.0)


class ExportAPITests(APITestCase):
    """Test case for exporting assignment and quiz scores as CSV files."""
    
    def setUp(self):
        self.instructor = PortalUser.objects.create(
            external_id='instructor_1',
            full_name='Instructor One',
            is_staff=True
        )
        self.other_instructor = PortalUser.objects.create(
            external_id='instructor_2',
            full_name='Instructor Two',
            is_staff=True
        )
        self.student = PortalUser.objects.create(
            external_id='student_1',
            full_name='Student One',
            email='student1@test.com'
        )
        
        self.course = CourseCache.objects.create(
            course_external_id=101,
            course_title='Test Course',
            course_code='TEST101'
        )
        
        # Instructor is assigned to course, but other_instructor is not.
        StaffAssignedCourse.objects.create(
            staff_external_id=self.instructor.external_id,
            course=self.course,
            role='instructor'
        )
        
        # Student is registered for course
        StudentRegisteredCourse.objects.create(
            student_external_id=self.student.external_id,
            course=self.course,
            session='2025/2026',
            semester='First Semester'
        )
        
        self.assignment = Assignment.objects.create(
            course=self.course,
            title='Test Assignment',
            description='Test Description',
            open_at=timezone.now(),
            due_at=timezone.now() + timezone.timedelta(days=1),
            max_marks=Decimal('10.00'),
            created_by=self.instructor
        )
        
        self.quiz = Quiz.objects.create(
            course=self.course,
            name='Test Quiz',
            max_grade=Decimal('100.00')
        )
        
        self.client = APIClient()

    def test_assignment_export_csv_authorized(self):
        """Test that an authorized instructor can download the assignment scores CSV."""
        # Create a submission for the student
        sub = AssignmentSubmission.objects.create(
            assignment=self.assignment,
            student_external_id=self.student.external_id,
            attempt_number=1,
            status='graded',
            marks=Decimal('8.50'),
            submitted_at=timezone.now(),
            graded_at=timezone.now()
        )
        
        self.client.force_authenticate(user=self.instructor)
        response = self.client.get(f'/api/staff/assessment/assignments/{self.assignment.id}/export/')
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response['Content-Type'], 'text/csv')
        self.assertIn(f"attachment; filename=\"assignment_Test_Assignment_scores.csv\"", response['Content-Disposition'])
        
        # Check CSV content
        content = response.content.decode('utf-8')
        lines = content.split('\r\n')
        self.assertTrue(len(lines) >= 2)
        self.assertIn('Student ID,Student Name,Student Email,Submission Status,Attempt Number,Submitted At,Score,Max Marks,Percentage', lines[0])
        self.assertIn('student_1,Student One,student1@test.com,Graded,1', content)
        self.assertIn('8.5,10.0,85.0', content)

    def test_assignment_export_csv_unauthorized(self):
        """Test that students or unassigned instructors are forbidden to download the assignment scores CSV."""
        # Test student forbidden
        self.client.force_authenticate(user=self.student)
        response = self.client.get(f'/api/staff/assessment/assignments/{self.assignment.id}/export/')
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        
        # Test unassigned instructor forbidden
        self.client.force_authenticate(user=self.other_instructor)
        response = self.client.get(f'/api/staff/assessment/assignments/{self.assignment.id}/export/')
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_quiz_export_csv_authorized(self):
        """Test that an authorized instructor can download the quiz scores CSV."""
        # Create a quiz attempt for the student
        attempt = QuizAttempt.objects.create(
            quiz=self.quiz,
            user_external_id=self.student.external_id,
            attempt_number=1,
            state='finished',
            started_at=timezone.now(),
            finished_at=timezone.now(),
            total_score=Decimal('90.00')
        )
        
        self.client.force_authenticate(user=self.instructor)
        response = self.client.get(f'/api/staff/assessment/quizzes/{self.quiz.id}/export/')
        
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response['Content-Type'], 'text/csv')
        self.assertIn(f"attachment; filename=\"quiz_Test_Quiz_scores.csv\"", response['Content-Disposition'])
        
        # Check CSV content
        content = response.content.decode('utf-8')
        lines = content.split('\r\n')
        self.assertTrue(len(lines) >= 2)
        self.assertIn('Student ID,Student Name,Student Email,Attempt State,Total Attempts,Best Attempt Number,Score,Max Grade,Percentage,Finished At', lines[0])
        self.assertIn('student_1,Student One,student1@test.com,Finished,1,1,90.0,100.0,90.0', content)

    def test_quiz_export_csv_unauthorized(self):
        """Test that students or unassigned instructors are forbidden to download the quiz scores CSV."""
        # Test student forbidden
        self.client.force_authenticate(user=self.student)
        response = self.client.get(f'/api/staff/assessment/quizzes/{self.quiz.id}/export/')
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

        # Test unassigned instructor forbidden
        self.client.force_authenticate(user=self.other_instructor)
        response = self.client.get(f'/api/staff/assessment/quizzes/{self.quiz.id}/export/')
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_quiz_excel_reports_include_class_scores_and_student_responses(self):
        category = QuestionCategory.objects.create(course=self.course, name='Quiz report bank')
        question = Question.objects.create(
            category=category,
            qtype='shortanswer',
            name='Math question',
            question_text='What is two plus two?',
            default_mark=Decimal('10.00'),
        )
        QuizQuestion.objects.create(quiz=self.quiz, question=question, order=1, max_mark=Decimal('10.00'))
        attempt = QuizAttempt.objects.create(
            quiz=self.quiz,
            user_external_id=self.student.external_id,
            attempt_number=1,
            state='finished',
            total_score=Decimal('90.00'),
            finished_at=timezone.now(),
        )
        QuestionAttempt.objects.create(
            quiz_attempt=attempt,
            question=question,
            display_order=1,
            response={'text': '<p><math><msqrt><mn>4</mn></msqrt></math></p>'},
            score=Decimal('10.00'),
        )

        self.client.force_authenticate(user=self.instructor)
        class_response = self.client.get(f'/api/staff/assessment/quizzes/{self.quiz.id}/report/')
        self.assertEqual(class_response.status_code, status.HTTP_200_OK)
        self.assertEqual(class_response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        class_book = load_workbook(BytesIO(class_response.content))
        self.assertEqual(class_book.sheetnames, ['Summary', 'Scores'])
        self.assertEqual(class_book['Scores']['A5'].value, self.student.external_id)
        self.assertEqual(class_book['Scores']['G5'].value, 90)

        student_response = self.client.get(f'/api/staff/assessment/quizzes/{self.quiz.id}/report/?attempt_id={attempt.id}')
        self.assertEqual(student_response.status_code, status.HTTP_200_OK)
        student_book = load_workbook(BytesIO(student_response.content))
        self.assertEqual(student_book.sheetnames, ['Student report', 'Responses'])
        self.assertEqual(student_book['Responses']['E5'].value, '4')

    def test_assessment_excel_report_includes_started_students(self):
        assessment = Assessment.objects.create(
            course=self.course,
            name='Assessment report',
            max_grade=Decimal('50.00'),
        )
        AssessmentAttempt.objects.create(
            assessment=assessment,
            user_external_id=self.student.external_id,
            attempt_number=1,
            state='in_progress',
            grade_scale=Decimal('50.00'),
        )

        self.client.force_authenticate(user=self.instructor)
        response = self.client.get(f'/api/staff/assessment/assessments/{assessment.id}/report/')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ['Summary', 'Scores'])
        self.assertEqual(workbook['Scores']['A5'].value, self.student.external_id)
        self.assertEqual(workbook['Scores']['D5'].value, 'In Progress')
